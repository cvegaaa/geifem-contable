"""Exportar a Excel, importar datos y administración de usuarios."""
import io

from openpyxl import Workbook, load_workbook

from tests.conftest import registrar

RANGO = {"desde": "2026-01-01", "hasta": "2026-12-31"}
XLSX = ("datos.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


async def _empresa(nuevo_cliente):
    c = nuevo_cliente()
    await registrar(c, "dueno@a.co", "900111222")
    return c


def _xlsx(filas) -> bytes:
    wb = Workbook()
    ws = wb.active
    for f in filas:
        ws.append(f)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


async def _importar(c, tipo, filas, simular=True, **params):
    return await c.post(f"/api/importar/{tipo}", params={"simular": str(simular).lower(), **params},
                        files={"archivo": (XLSX[0], _xlsx(filas), XLSX[1])})


# ------------------------------------------------------------------ exportar

async def test_exportar_tabla_a_excel(nuevo_cliente):
    c = await _empresa(nuevo_cliente)
    r = await c.post("/api/exportar/xlsx", json={
        "titulo": "Balance de prueba", "columnas": [{"key": "codigo", "label": "Código"}, {"key": "saldo", "label": "Saldo"}],
        "filas": [{"codigo": "1105", "saldo": 1500.5}, {"codigo": "2205", "saldo": -300}],
        "resumen": {"debitos": 1500.5}})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    valores = [[c.value for c in fila] for fila in ws.iter_rows()]
    assert "Balance de prueba" in valores[0][0] and valores[5] == ["1105", 1500.5] and ["2205", -300] in valores


# ------------------------------------------------------------------ importar

async def test_plantillas_se_descargan_y_se_leen_de_vuelta(nuevo_cliente):
    c = await _empresa(nuevo_cliente)
    for tipo in ("terceros", "productos", "saldos-iniciales"):
        r = await c.get(f"/api/importar/{tipo}/plantilla")
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames[1] == "Instrucciones" and wb.active["A2"].value is not None
    # La propia plantilla de terceros es un archivo importable (la fila 2 es el ejemplo).
    plantilla = (await c.get("/api/importar/terceros/plantilla")).content
    r = await c.post("/api/importar/terceros", params={"simular": "false"}, files={"archivo": XLSX[:1] + (plantilla, XLSX[1])})
    assert r.status_code == 200 and r.json()["importado"] and r.json()["creados"] == 1


async def test_importar_terceros_valida_todo_o_nada(nuevo_cliente):
    c = await _empresa(nuevo_cliente)
    cab = ["Nombre *", "Tipo", "Número de documento *", "Régimen"]
    buenas = [cab, ["Ana Ruiz", "cliente", "1001", "ORDINARIO"], ["Beta SAS", "proveedor", "900200300", "SIMPLE"]]
    r = await _importar(c, "terceros", buenas)
    assert r.json()["validas"] == 2 and not r.json()["errores"] and not r.json()["importado"]   # simulación: no guarda
    assert (await c.get("/api/configuracion/terceros")).json() == []

    malas = [cab, ["Ana Ruiz", "cliente", "1001", "ORDINARIO"], ["Repetida", "cliente", "1001", "ORDINARIO"],
             ["Sin régimen válido", "cliente", "1002", "OTRO"], ["", "cliente", "1003", "ORDINARIO"]]
    r = await _importar(c, "terceros", malas, simular=False)
    j = r.json()
    assert not j["importado"] and [e["fila"] for e in j["errores"]] == [3, 4, 5]
    assert (await c.get("/api/configuracion/terceros")).json() == []      # con errores no se guarda nada

    r = await _importar(c, "terceros", buenas, simular=False)
    assert r.json()["importado"] and r.json()["creados"] == 2
    terceros = {t["numero_documento"]: t for t in (await c.get("/api/configuracion/terceros")).json()}
    assert terceros["900200300"]["regimen"] == "SIMPLE" and terceros["1001"]["responsable_iva"] is True
    # Reimportar los mismos documentos se rechaza.
    assert (await _importar(c, "terceros", buenas)).json()["errores"]


async def test_importar_productos_respeta_el_perfil_tributario(nuevo_cliente):
    c = await _empresa(nuevo_cliente)
    cab = ["Código (SKU) *", "Nombre *", "Precio de venta", "Tarifa de IVA", "Controla inventario"]
    r = await _importar(c, "productos", [cab, ["A-1", "Camisa", 45000, 19, "SI"], ["A-2", "Servicio", 10000, 7, "NO"]])
    assert [e["fila"] for e in r.json()["errores"]] == [3]
    r = await _importar(c, "productos", [cab, ["A-1", "Camisa", 45000, 19, "SI"], ["A-3", "Servicio", 10000, 0, "NO"]], simular=False)
    assert r.json()["creados"] == 2
    prods = {p["sku"]: p for p in (await c.get("/api/inventario/catalogo-productos")).json()}
    assert prods["A-3"]["inventariable"] is False and prods["A-1"]["tarifa_iva"] == 19


async def test_importar_saldos_iniciales_crea_apertura_cuadrada(nuevo_cliente):
    c = await _empresa(nuevo_cliente)
    cab = ["Código de cuenta *", "Descripción", "Débito", "Crédito"]
    desc = [cab, ["1110", "Bancos", 5_000_000, None], ["3105", "Capital", None, 4_000_000]]
    r = await _importar(c, "saldos-iniciales", desc, fecha="2026-01-01")
    assert r.json()["errores"][0]["mensaje"].startswith("los saldos no cuadran")

    ok = [cab, ["1110", "Bancos", 5_000_000, None], ["3105", "Capital", None, 5_000_000]]
    r = await _importar(c, "saldos-iniciales", ok, simular=False, fecha="2026-01-01")
    assert r.json()["importado"] and r.json()["detalle"]["comprobante"] == "AP-1"
    bal = (await c.get("/api/reportes/contables/balance-de-prueba", params=RANGO)).json()
    assert {f["codigo"]: f["saldo"] for f in bal["filas"]} == {"1110": 5_000_000, "3105": -5_000_000}

    r = await _importar(c, "saldos-iniciales", [cab, ["999999", "X", 1, None]], fecha="2026-01-01")
    assert "no existe" in r.json()["errores"][0]["mensaje"]
    assert (await c.post("/api/importar/saldos-iniciales", files={"archivo": XLSX[:1] + (_xlsx(ok), XLSX[1])})).status_code == 422


async def test_importar_acepta_csv_y_rechaza_archivos_malos(nuevo_cliente):
    c = await _empresa(nuevo_cliente)
    csv = "Nombre *;Número de documento *\nAna;1001\nLuis;1002\n".encode("utf-8")
    r = await c.post("/api/importar/terceros", params={"simular": "false"}, files={"archivo": ("t.csv", csv, "text/csv")})
    assert r.json()["creados"] == 2
    r = await c.post("/api/importar/terceros", files={"archivo": ("x.xlsx", b"no es excel", XLSX[1])})
    assert r.status_code == 422
    r = await c.post("/api/importar/terceros", files={"archivo": ("t.csv", "Otra;Cosa\n1;2\n".encode(), "text/csv")})
    assert r.status_code == 422 and "Faltan columnas" in r.json()["detail"]


# ------------------------------------------------------------------ usuarios

async def _con_usuario(c, nuevo_cliente, email="cont@a.co", rol="contador"):
    r = await c.post("/api/configuracion/usuarios-roles", json={
        "nombre": "Usuario", "email": email, "password": "clave-segura-1", "rol": rol})
    assert r.status_code == 201, r.text
    otro = nuevo_cliente()
    await otro.post("/api/auth/login", json={"email": email, "password": "clave-segura-1"})
    otro.headers["X-Empresa-Id"] = c.headers["X-Empresa-Id"]
    return r.json(), otro


async def test_permisos_dependen_del_rol_y_se_editan_en_roles(nuevo_cliente):
    c = await _empresa(nuevo_cliente)
    u, sesion = await _con_usuario(c, nuevo_cliente)
    assert (await sesion.get("/api/nomina/empleados")).status_code == 200
    assert u["permisos"]["nomina"] and u["rol_nombre"] == "Contador"

    # El acceso se edita en el rol y aplica de inmediato a sus usuarios, sin cerrarles la sesión.
    catalogo = (await c.get("/api/configuracion/roles")).json()
    contador = next(r for r in catalogo["roles"] if r["clave"] == "contador")
    assert contador["usuarios"] == 1 and "aprobar" in catalogo["acciones"]
    r = await c.put("/api/configuracion/roles/contador", json={
        "nombre": "Contador", "permisos": {**contador["permisos"], "nomina": []}, "requiere_aprobacion": False})
    assert r.status_code == 200 and r.json()["permisos"]["nomina"] == []
    assert (await sesion.get("/api/nomina/empleados")).status_code == 403
    assert (await sesion.get("/api/auth/me")).json()["permisos"]["nomina"] == []

    # Los roles propios se crean, se asignan y no se borran con usuarios.
    r = await c.post("/api/configuracion/roles", json={
        "nombre": "Cajero", "permisos": {"facturacion": ["leer", "crear"], "tesoreria": ["leer"]}, "requiere_aprobacion": True})
    assert r.status_code == 201 and r.json()["clave"] == "cajero" and r.json()["permisos"]["nomina"] == []
    assert (await c.post("/api/configuracion/roles", json={"nombre": "Cajero", "permisos": {}})).status_code == 409
    r = await c.put(f"/api/configuracion/usuarios-roles/{u['id']}", json={"nombre": "U", "rol": "cajero", "activo": True})
    assert r.status_code == 200 and r.json()["rol_nombre"] == "Cajero" and r.json()["aprobacion_requerida"] is True
    assert r.json()["permisos"]["facturacion"] == ["leer", "crear"]
    assert (await c.delete("/api/configuracion/roles/cajero")).status_code == 409
    assert (await c.delete("/api/configuracion/roles/contador")).status_code == 409           # rol base
    assert (await c.put(f"/api/configuracion/usuarios-roles/{u['id']}", json={"nombre": "U", "rol": "inventado", "activo": True})).status_code == 422
    assert (await c.put("/api/configuracion/roles/admin", json={"nombre": "xx", "permisos": {}})).status_code == 409
    assert (await c.post("/api/configuracion/roles", json={"nombre": "Raro", "permisos": {"magia": ["leer"]}})).status_code == 422
    # Un usuario sin ser administrador no cambia roles.
    nueva = nuevo_cliente()
    await nueva.post("/api/auth/login", json={"email": "cont@a.co", "password": "clave-segura-1"})
    nueva.headers["X-Empresa-Id"] = c.headers["X-Empresa-Id"]
    assert (await nueva.put("/api/configuracion/roles/cajero", json={"nombre": "Cajero", "permisos": {}})).status_code == 403


async def test_no_se_puede_dejar_la_empresa_sin_administrador(nuevo_cliente):
    c = await _empresa(nuevo_cliente)
    yo = (await c.get("/api/auth/me")).json()
    r = await c.put(f"/api/configuracion/usuarios-roles/{yo['id']}", json={"nombre": "Yo", "rol": "contador", "activo": True})
    assert r.status_code == 409
    r = await c.put(f"/api/configuracion/usuarios-roles/{yo['id']}", json={"nombre": "Yo", "rol": "admin", "activo": False})
    assert r.status_code == 409
    # Un no administrador no gestiona usuarios.
    u, sesion = await _con_usuario(c, nuevo_cliente)
    r = await sesion.put(f"/api/configuracion/usuarios-roles/{yo['id']}", json={"nombre": "X", "rol": "admin", "activo": True})
    assert r.status_code == 403


async def test_restablecer_password_y_aislamiento(nuevo_cliente):
    c = await _empresa(nuevo_cliente)
    u, sesion = await _con_usuario(c, nuevo_cliente)
    r = await c.post(f"/api/configuracion/usuarios-roles/{u['id']}/restablecer-password", json={"nueva": "corta"})
    assert r.status_code == 422
    r = await c.post(f"/api/configuracion/usuarios-roles/{u['id']}/restablecer-password", json={"nueva": "nueva-clave-77"})
    assert r.status_code == 204
    assert (await sesion.get("/api/auth/me")).status_code == 401
    ok = await nuevo_cliente().post("/api/auth/login", json={"email": "cont@a.co", "password": "nueva-clave-77"})
    assert ok.status_code == 200

    # Un administrador de otra empresa no puede tocar a este usuario.
    otra = nuevo_cliente()
    await registrar(otra, "otro@b.co", "800999888", "Otra SAS")
    r = await otra.put(f"/api/configuracion/usuarios-roles/{u['id']}", json={"nombre": "X", "rol": "admin", "activo": True})
    assert r.status_code == 404
    assert (await otra.post(f"/api/configuracion/usuarios-roles/{u['id']}/restablecer-password", json={"nueva": "nueva-clave-88"})).status_code == 404
