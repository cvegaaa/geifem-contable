"""Archivos: exportar tablas a Excel e importar datos desde Excel o CSV.

- Exportar: cualquier tabla que el usuario ve (reportes y listados) se descarga como .xlsx.
- Importar: terceros, productos y saldos iniciales, con plantilla descargable, validación previa
  (`simular=true`) y carga todo-o-nada: si hay errores no se guarda nada.
"""
import csv
import io
import re
import unicodedata
from datetime import date, datetime

from bson import ObjectId
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from api.contabilidad import DOC_TIPO, PREFIJOS, _aplicar
from api.inventario import _validar_cuentas_producto
from core.auth.dependencies import get_empresa_activa, require_permiso
from core import nit
from core.dinero import centavos
from core.documentos import (
    ahora,
    auditar,
    en_transaccion,
    formatear_numero,
    serializar,
    siguiente_consecutivo,
)
from db import get_db

router = APIRouter(prefix="/api", tags=["archivos"])

MAX_FILAS_EXPORTAR = 50_000
MAX_BYTES_IMPORTAR = 5 * 1024 * 1024
MAX_FILAS_IMPORTAR = 5_000
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------------------
# Exportar
# ---------------------------------------------------------------------------

class ColumnaExport(BaseModel):
    key: str
    label: str


class ExportarEntrada(BaseModel):
    titulo: str = Field(min_length=1, max_length=120)
    columnas: list[ColumnaExport] = Field(min_length=1, max_length=60)
    filas: list[dict] = Field(max_length=MAX_FILAS_EXPORTAR)
    resumen: dict | None = None


def _celda(valor):
    """Valores simples para Excel: números como números, el resto como texto."""
    if valor is None or isinstance(valor, (int, float, bool)):
        return valor
    if isinstance(valor, (dict, list)):
        return str(valor)
    return str(valor)


def _libro(titulo: str, columnas: list[tuple[str, str]], filas: list[dict], pie: dict | None = None,
           encabezado: list[str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]\*\?/\\:]", " ", titulo)[:31] or "Datos"
    fila = 1
    for linea in encabezado or []:
        ws.cell(row=fila, column=1, value=linea).font = Font(bold=fila == 1, size=13 if fila == 1 else 10)
        fila += 1
    if encabezado:
        fila += 1
    cab = fila
    for i, (_, etiqueta) in enumerate(columnas, start=1):
        c = ws.cell(row=cab, column=i, value=etiqueta)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="14315E")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, datos in enumerate(filas, start=cab + 1):
        for i, (clave, _) in enumerate(columnas, start=1):
            c = ws.cell(row=r, column=i, value=_celda(datos.get(clave)))
            if isinstance(c.value, float) and not float(c.value).is_integer():
                c.number_format = "#,##0.00"
            elif isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                c.number_format = "#,##0"
    fin = cab + len(filas)
    if pie:
        fin += 2
        for clave, valor in pie.items():
            ws.cell(row=fin, column=1, value=clave.replace("_", " ").capitalize()).font = Font(bold=True)
            ws.cell(row=fin, column=2, value=_celda(valor))
            fin += 1
    for i, (clave, etiqueta) in enumerate(columnas, start=1):
        ancho = max([len(str(etiqueta))] + [len(str(f.get(clave, "") or "")) for f in filas[:300]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(ancho, 10), 60)
    ws.freeze_panes = ws.cell(row=cab + 1, column=1)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _respuesta_xlsx(contenido: bytes, nombre: str) -> StreamingResponse:
    seguro = re.sub(r"[^A-Za-z0-9_.-]+", "_", nombre) or "datos"
    return StreamingResponse(
        io.BytesIO(contenido), media_type=XLSX,
        headers={"Content-Disposition": f'attachment; filename="{seguro}.xlsx"'})


@router.post("/exportar/xlsx")
async def exportar_xlsx(
    payload: ExportarEntrada,
    empresa_id: str = Depends(get_empresa_activa),
    _=Depends(require_permiso("reportes", "leer")),
):
    """Convierte en Excel una tabla que el cliente ya tiene en pantalla."""
    empresa = await get_db().empresas.find_one({"_id": ObjectId(empresa_id)}) or {}
    encabezado = [payload.titulo, f"{empresa.get('razon_social', '')} · NIT {empresa.get('nit', '')}",
                  f"Generado el {datetime.now().strftime('%Y-%m-%d %H:%M')}"]
    contenido = _libro(payload.titulo, [(c.key, c.label) for c in payload.columnas], payload.filas,
                       payload.resumen, encabezado)
    return _respuesta_xlsx(contenido, f"{payload.titulo}-{date.today().isoformat()}")


# ---------------------------------------------------------------------------
# Importar: definición de cada tipo
# ---------------------------------------------------------------------------

def _norm(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", t.lower()).strip("_")


def _bool(v, defecto=None):
    if v is None or str(v).strip() == "":
        return defecto
    t = _norm(v)
    if t in ("si", "s", "1", "true", "x", "yes", "y"):
        return True
    if t in ("no", "n", "0", "false"):
        return False
    raise ValueError(f"valor «{v}» no válido (usa SI o NO)")


def _num(v, defecto=None):
    if v is None or str(v).strip() == "":
        return defecto
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace("$", "").replace(" ", "")
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        raise ValueError(f"«{v}» no es un número")


def _texto(v):
    return str(v).strip() if v is not None and str(v).strip() != "" else None


def _fecha(v):
    if isinstance(v, (datetime, date)):
        return v.date().isoformat() if isinstance(v, datetime) else v.isoformat()
    t = _texto(v)
    if not t:
        return None
    try:
        return date.fromisoformat(t[:10]).isoformat()
    except ValueError:
        raise ValueError(f"fecha «{v}» no válida (usa AAAA-MM-DD)")


PLANTILLAS = {
    "terceros": {
        "titulo": "Terceros",
        "columnas": [
            ("nombre", "Nombre *"), ("tipo", "Tipo (cliente, proveedor, ambos, otro)"), ("tipo_persona", "Persona (natural, juridica)"),
            ("tipo_documento", "Tipo de documento (CC, NIT, CE, PA, TI)"), ("numero_documento", "Número de documento *"), ("dv", "DV"),
            ("direccion", "Dirección"), ("ciudad", "Ciudad"), ("email", "Email"), ("telefono", "Teléfono"),
            ("regimen", "Régimen (ORDINARIO, SIMPLE)"), ("responsable_iva", "Responsable de IVA (SI/NO)"),
            ("declarante_renta", "Declarante de renta (SI/NO)"), ("autorretenedor", "Autorretenedor (SI/NO)"),
            ("agente_retenedor", "Agente retenedor (SI/NO)"), ("agente_reteiva", "Agente de ReteIVA (SI/NO)"),
            ("gran_contribuyente", "Gran contribuyente (SI/NO)"),
        ],
        "ejemplo": ["Distribuidora Andina SAS", "proveedor", "juridica", "NIT", "900123456", "8", "Cra 10 # 20-30", "Bogotá",
                    "compras@andina.co", "3001234567", "ORDINARIO", "SI", "SI", "NO", "NO", "NO", "NO"],
        "ayuda": ["Los campos con * son obligatorios. El número de documento no puede repetirse.",
                  "Las columnas de perfil tributario definen las retenciones: déjalas vacías si no aplican."],
    },
    "productos": {
        "titulo": "Productos",
        "columnas": [
            ("sku", "Código (SKU) *"), ("nombre", "Nombre *"), ("categoria", "Categoría"), ("precio_venta", "Precio de venta"),
            ("tipo_iva", "Tipo de IVA (gravado, excluido, exento)"), ("tarifa_iva", "Tarifa de IVA (0, 5, 19)"),
            ("inventariable", "Controla inventario (SI/NO)"),
        ],
        "ejemplo": ["CAM-001", "Camiseta básica", "Ropa", 45000, "gravado", 19, "SI"],
        "ayuda": ["El código (SKU) no puede repetirse. Si tu empresa no es responsable de IVA, deja la tarifa en 0.",
                  "Las existencias iniciales se cargan con una entrada de mercancía o un ajuste, no desde aquí."],
    },
    "saldos-iniciales": {
        "titulo": "Saldos iniciales",
        "columnas": [
            ("cuenta", "Código de cuenta *"), ("descripcion", "Descripción"), ("debito", "Débito"), ("credito", "Crédito"),
            ("tercero_documento", "Documento del tercero (opcional)"),
            ("documento", "Número del documento (cartera)"), ("fecha_documento", "Fecha del documento (cartera)"),
            ("vencimiento", "Vencimiento (cartera)"),
        ],
        "ejemplo": ["1110", "Saldo en bancos", 5000000, None, None, None, None, None],
        "ayuda": ["Genera un comprobante de apertura contabilizado. Débitos y créditos deben cuadrar.",
                  "Indica la fecha de apertura al importar. Cada línea lleva débito o crédito, no ambos.",
                  "CARTERA: las líneas de la cuenta de clientes (débito) y de proveedores (crédito) deben llevar el documento del tercero y el "
                  "número de la factura; así quedan como cuentas por cobrar o por pagar por documento, con su fecha y vencimiento, y aparecen en "
                  "Tesorería, en la antigüedad de cartera y en los reportes. Una línea por factura pendiente."],
    },
}


@router.get("/importar/{tipo}/plantilla")
async def descargar_plantilla(tipo: str, _=Depends(require_permiso("configuracion", "leer"))):
    p = PLANTILLAS.get(tipo)
    if not p:
        raise HTTPException(404, "Tipo de importación desconocido")
    wb = Workbook()
    ws = wb.active
    ws.title = p["titulo"]
    for i, (_, etiqueta) in enumerate(p["columnas"], start=1):
        c = ws.cell(row=1, column=i, value=etiqueta)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="14315E")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = max(len(etiqueta) * 0.9, 16)
        ws.cell(row=2, column=i, value=p["ejemplo"][i - 1])
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "A2"
    ayuda = wb.create_sheet("Instrucciones")
    ayuda.column_dimensions["A"].width = 110
    ayuda.cell(row=1, column=1, value=f"Plantilla de {p['titulo'].lower()}").font = Font(bold=True, size=13)
    for i, linea in enumerate(["La fila 2 es un ejemplo: bórrala o reemplázala.", *p["ayuda"]], start=3):
        ayuda.cell(row=i, column=1, value=linea)
    salida = io.BytesIO()
    wb.save(salida)
    return _respuesta_xlsx(salida.getvalue(), f"plantilla-{tipo}")


def _leer_filas(nombre: str, contenido: bytes, columnas: list[tuple[str, str]]) -> list[tuple[int, dict]]:
    """Filas del archivo como {clave: valor}, tolerando etiquetas con o sin sufijos de ayuda."""
    if nombre.lower().endswith(".csv"):
        texto = contenido.decode("utf-8-sig", errors="replace")
        dialecto = ";" if texto.splitlines()[0].count(";") > texto.splitlines()[0].count(",") else ","
        filas = [list(f) for f in csv.reader(io.StringIO(texto), delimiter=dialecto)]
    else:
        try:
            wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(422, "No se pudo leer el archivo: debe ser .xlsx o .csv")
        filas = [list(f) for f in wb.worksheets[0].iter_rows(values_only=True)]
    if not filas:
        raise HTTPException(422, "El archivo está vacío")
    if len(filas) - 1 > MAX_FILAS_IMPORTAR:
        raise HTTPException(422, f"El archivo supera {MAX_FILAS_IMPORTAR} filas: divídelo en partes")

    # Cada encabezado se asocia por su nombre exacto (sin lo que va entre paréntesis) o por su clave.
    por_nombre = {_norm(e.split("(")[0].replace("*", "")): c for c, e in columnas}
    por_nombre.update({c: c for c, _ in columnas})
    indice: dict[int, str] = {}
    for i, cab in enumerate(filas[0]):
        base = _norm(str(cab or "").split("(")[0].replace("*", ""))
        if base in por_nombre and por_nombre[base] not in indice.values():
            indice[i] = por_nombre[base]
    faltan = [e for c, e in columnas if "*" in e and c not in indice.values()]
    if faltan:
        raise HTTPException(422, f"Faltan columnas obligatorias: {', '.join(f.replace(' *', '') for f in faltan)}")
    salida = []
    for n, fila in enumerate(filas[1:], start=2):
        if not any(v not in (None, "") for v in fila):
            continue
        salida.append((n, {clave: (fila[i] if i < len(fila) else None) for i, clave in indice.items()}))
    return salida


class Resultado(BaseModel):
    tipo: str
    simulacion: bool
    total: int
    validas: int
    errores: list[dict]
    importado: bool
    creados: int = 0
    detalle: dict | None = None


@router.post("/importar/{tipo}", response_model=Resultado)
async def importar(
    tipo: str,
    archivo: UploadFile = File(...),
    simular: bool = Query(True, description="true = solo validar, sin guardar"),
    fecha: str | None = Query(None, description="Fecha de apertura (solo saldos iniciales)"),
    empresa_id: str = Depends(get_empresa_activa),
    usuario: dict = Depends(require_permiso("configuracion", "crear")),
):
    if tipo not in PLANTILLAS:
        raise HTTPException(404, "Tipo de importación desconocido")
    contenido = await archivo.read(MAX_BYTES_IMPORTAR + 1)
    if len(contenido) > MAX_BYTES_IMPORTAR:
        raise HTTPException(413, "El archivo supera 5 MB")
    filas = _leer_filas(archivo.filename or "", contenido, PLANTILLAS[tipo]["columnas"])
    if not filas:
        raise HTTPException(422, "El archivo no tiene filas de datos")

    if tipo == "terceros":
        docs, errores = await _validar_terceros(empresa_id, filas)
        return await _guardar_lote("terceros", docs, errores, len(filas), simular, empresa_id, usuario)
    if tipo == "productos":
        docs, errores = await _validar_productos(empresa_id, filas)
        return await _guardar_lote("productos", docs, errores, len(filas), simular, empresa_id, usuario)
    return await _importar_saldos(empresa_id, usuario, filas, simular, fecha)


async def _validar_terceros(empresa_id: str, filas) -> tuple[list[dict], list[dict]]:
    existentes = {t["numero_documento"] async for t in get_db().terceros.find(
        {"empresa_id": empresa_id}, {"numero_documento": 1})}
    vistos, docs, errores = set(), [], []
    for n, f in filas:
        try:
            nombre, doc = _texto(f.get("nombre")), _texto(f.get("numero_documento"))
            if not nombre or not doc:
                raise ValueError("faltan el nombre o el número de documento")
            doc = doc[:-2] if doc.endswith(".0") else doc
            if doc in existentes or doc in vistos:
                raise ValueError(f"el documento {doc} ya existe o está repetido en el archivo")
            vistos.add(doc)
            tipo = (_norm(f.get("tipo")) or "cliente")
            if tipo not in ("cliente", "proveedor", "ambos", "otro"):
                raise ValueError("tipo debe ser cliente, proveedor, ambos u otro")
            regimen = (_texto(f.get("regimen")) or "ORDINARIO").upper()
            if regimen not in ("ORDINARIO", "SIMPLE"):
                raise ValueError("régimen debe ser ORDINARIO o SIMPLE")
            tipo_doc = (_texto(f.get("tipo_documento")) or "CC").upper()
            dv = _texto(f.get("dv"))
            if tipo_doc == "NIT":
                esperado = nit.digito_verificacion(doc)
                if esperado is None:
                    raise ValueError("el NIT debe ser numérico (máximo 15 dígitos)")
                if dv and dv.removesuffix(".0") != esperado:
                    raise ValueError(f"el DV de {doc} no corresponde: debería ser {esperado}")
                doc, dv = nit.limpiar(doc), esperado
            docs.append({
                "nombre": nombre, "tipo": tipo, "tipo_persona": _norm(f.get("tipo_persona")) or "natural",
                "tipo_documento": tipo_doc, "numero_documento": doc,
                "dv": dv, "direccion": _texto(f.get("direccion")), "ciudad": _texto(f.get("ciudad")),
                "email": _texto(f.get("email")), "telefono": _texto(f.get("telefono")), "regimen": regimen,
                "responsable_iva": _bool(f.get("responsable_iva"), True), "declarante_renta": _bool(f.get("declarante_renta"), True),
                "autorretenedor": _bool(f.get("autorretenedor"), False), "agente_retenedor": _bool(f.get("agente_retenedor"), False),
                "agente_reteiva": _bool(f.get("agente_reteiva"), False), "gran_contribuyente": _bool(f.get("gran_contribuyente"), False),
            })
        except ValueError as exc:
            errores.append({"fila": n, "mensaje": str(exc)})
    return docs, errores


async def _validar_productos(empresa_id: str, filas) -> tuple[list[dict], list[dict]]:
    existentes = {p["sku"] async for p in get_db().productos.find({"empresa_id": empresa_id}, {"sku": 1})}
    vistos, docs, errores = set(), [], []
    for n, f in filas:
        try:
            sku, nombre = _texto(f.get("sku")), _texto(f.get("nombre"))
            if not sku or not nombre:
                raise ValueError("faltan el código o el nombre")
            sku = sku[:-2] if sku.endswith(".0") else sku
            if sku in existentes or sku in vistos:
                raise ValueError(f"el código {sku} ya existe o está repetido en el archivo")
            vistos.add(sku)
            tarifa = _num(f.get("tarifa_iva"), 0)
            if tarifa not in (0, 5, 19):
                raise ValueError("la tarifa de IVA debe ser 0, 5 o 19")
            doc = {"sku": sku, "nombre": nombre, "categoria": _texto(f.get("categoria")),
                   "precio_venta": _num(f.get("precio_venta"), 0),
                   "tipo_iva": _norm(f.get("tipo_iva")) or "gravado", "tarifa_iva": tarifa,
                   "inventariable": _bool(f.get("inventariable"), True)}
            try:
                await _validar_cuentas_producto(doc, empresa_id)
            except HTTPException as exc:
                raise ValueError(str(exc.detail))
            docs.append(doc)
        except ValueError as exc:
            errores.append({"fila": n, "mensaje": str(exc)})
    return docs, errores


async def _guardar_lote(coleccion: str, docs: list[dict], errores: list[dict], total: int, simular: bool,
                        empresa_id: str, usuario: dict) -> Resultado:
    importado = False
    if not simular and not errores and docs:
        momento = ahora()
        await get_db()[coleccion].insert_many([
            {**d, "empresa_id": empresa_id, "activo": True, "importado": True,
             "creado_por": usuario["id"], "fecha_creacion": momento} for d in docs])
        importado = True
    return Resultado(tipo=coleccion, simulacion=simular, total=total, validas=len(docs), errores=errores,
                     importado=importado, creados=len(docs) if importado else 0)


async def _importar_saldos(empresa_id: str, usuario: dict, filas, simular: bool, fecha: str | None) -> Resultado:
    db = get_db()
    errores: list[dict] = []
    if not fecha:
        raise HTTPException(422, "Indica la fecha de apertura")
    try:
        date.fromisoformat(fecha)
    except ValueError:
        raise HTTPException(422, "Fecha de apertura inválida (usa AAAA-MM-DD)")
    cuentas = {c["codigo"]: c async for c in db.plan_cuentas.find({"empresa_id": empresa_id})}
    terceros = {t["numero_documento"]: str(t["_id"]) async for t in db.terceros.find({"empresa_id": empresa_id})}
    lineas = []
    for n, f in filas:
        try:
            codigo = _texto(f.get("cuenta"))
            codigo = codigo[:-2] if codigo and codigo.endswith(".0") else codigo
            cuenta = cuentas.get(codigo or "")
            if not cuenta or cuenta.get("activo") is False:
                raise ValueError(f"la cuenta {codigo} no existe en el plan de cuentas")
            if cuenta.get("acepta_movimientos") is False:
                raise ValueError(f"la cuenta {codigo} es de agrupación y no recibe movimientos")
            deb, cre = _num(f.get("debito"), 0), _num(f.get("credito"), 0)
            if bool(deb) == bool(cre):
                raise ValueError("cada línea debe tener valor en débito o en crédito, no en ambos")
            if deb < 0 or cre < 0:
                raise ValueError("los valores no pueden ser negativos")
            doc = _texto(f.get("tercero_documento"))
            tercero_id = None
            if doc:
                tercero_id = terceros.get(doc[:-2] if doc.endswith(".0") else doc)
                if not tercero_id:
                    raise ValueError(f"el tercero {doc} no existe: impórtalo primero")
            documento = _texto(f.get("documento"))
            lineas.append({"cuenta_puc_id": str(cuenta["_id"]), "debito": deb, "credito": cre,
                           "tercero_id": tercero_id, "descripcion": _texto(f.get("descripcion")),
                           "documento_numero": documento[:-2] if documento and documento.endswith(".0") else documento,
                           "fecha_documento": _fecha(f.get("fecha_documento")),
                           "fecha_vencimiento": _fecha(f.get("vencimiento"))})
        except ValueError as exc:
            errores.append({"fila": n, "mensaje": str(exc)})
    total_deb = sum(centavos(l["debito"]) for l in lineas)
    total_cre = sum(centavos(l["credito"]) for l in lineas)
    if not errores and total_deb != total_cre:
        errores.append({"fila": 0, "mensaje": f"los saldos no cuadran: débitos {total_deb / 100:,.2f} y créditos {total_cre / 100:,.2f}"})
    detalle = {"debitos": total_deb / 100, "creditos": total_cre / 100}
    if simular or errores or not lineas:
        return Resultado(tipo="saldos-iniciales", simulacion=simular, total=len(filas), validas=len(lineas),
                         errores=errores, importado=False, detalle=detalle)

    async def _tx(s):
        consecutivo = await siguiente_consecutivo(empresa_id, "comprobante:apertura", s)
        doc = {"tipo": "apertura", "fecha": fecha, "estado": "contabilizado", "descripcion": "Saldos iniciales importados",
               "tercero_id": None, "lineas": lineas, "empresa_id": empresa_id, "prefijo": PREFIJOS["apertura"],
               "consecutivo": consecutivo, "numero": formatear_numero(PREFIJOS["apertura"], consecutivo),
               "creado_por": usuario["id"], "fecha_creacion": ahora(), "importado": True}
        res = await db.comprobantes_contables.insert_one(doc, session=s)
        await _aplicar(empresa_id, str(res.inserted_id), doc, s)
        await auditar(empresa_id, DOC_TIPO, str(res.inserted_id), doc["numero"], usuario["id"], "creado", s,
                      {"origen": "importación"})
        return res.inserted_id

    nuevo = await en_transaccion(_tx)
    comp = serializar(await db.comprobantes_contables.find_one({"_id": nuevo}))
    return Resultado(tipo="saldos-iniciales", simulacion=False, total=len(filas), validas=len(lineas), errores=[],
                     importado=True, creados=1, detalle={**detalle, "comprobante": comp["numero"]})
